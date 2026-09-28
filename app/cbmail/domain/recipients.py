"""Извлечение адресов из загруженного файла со списком получателей.

Списки к рассылке почти никогда не приходят текстом в поле: их присылают
выгрузкой из 1С, CRM или просто «файлом из почты» — .xlsx, .xls, реже .csv.
Раньше такой файл приходилось открывать, копировать колонку и вставлять
руками; на списке в несколько тысяч строк это и долго, и ошибочно.

Поддерживаются четыре формата, и каждый разбирается по-своему:

* **TXT и CSV** — текст целиком, адреса выбираются по разделителям. Кодировка
  определяется по содержимому: UTF-8 (с BOM или без), иначе cp1251 — в ней
  приходит большинство выгрузок из старых российских систем.
* **XLSX** — через openpyxl, по всем листам и ячейкам. Адрес может лежать не в
  первой колонке и не под заголовком «email», поэтому просматривается весь
  лист, а не угаданный столбец.
* **XLS (Excel 97–2003)** — через xlrd. Если библиотека не справилась (битый
  или нестандартный файл), включается запасной путь: побайтовый поиск адресов
  в содержимом. Формат старый, файлы к нему прилагаются соответствующие, и
  вернуть человеку «не смог прочитать» там, где адреса видны глазами, — плохой
  ответ.

Для двоичного поиска используется строгая регулярка: в тексте разделителем
служит пробел или запятая и достаточно «всё, что не разделитель», а внутри
.xls адрес окружён произвольными байтами, и нестрогое правило приклеивает их
к имени — из служебной сигнатуры файла получается «ÐÏàivan@example.ru».
"""

from __future__ import annotations

import io
import logging
import re
from dataclasses import dataclass

from cbmail.domain.emails import is_valid_email, normalize_email

logger = logging.getLogger(__name__)

#: Предел размера файла. Больше — это не список рассылки, а чья-то база.
MAX_FILE_BYTES = 10 * 1024 * 1024

#: Что принимаем. Расширение проверяется до чтения: так понятнее отказ.
ALLOWED_EXTENSIONS = (".txt", ".csv", ".xlsx", ".xls")

#: Сигнатура составного документа OLE2 — контейнера старых файлов Excel.
_OLE2_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"

#: Разбор текста: разделителями служат пробелы, запятые, точки с запятой и скобки.
_TEXT_EMAIL_RE = re.compile(r"[^\s,;<>()\"']+@[^\s,;<>()\"']+\.[^\s,;<>()\"']+")

#: Разбор двоичного содержимого: перечислены допустимые символы, а не запрещённые.
_BINARY_EMAIL_RE = re.compile(
    r"[A-Za-z0-9._%+-]+@[A-Za-z0-9-]+(?:\.[A-Za-z0-9-]+)*\.[A-Za-z]{2,}"
)


class RecipientFileError(Exception):
    """Файл нельзя разобрать — текст исключения показывается человеку."""


@dataclass(frozen=True, slots=True)
class ImportResult:
    """Что получилось из файла.

    ``skipped_invalid`` и ``skipped_duplicates`` показываются в интерфейсе:
    «загружено 1240 адресов» без упоминания о 60 отброшенных выглядит как
    потеря данных, и человек идёт искать, куда делись строки.
    """

    emails: list[str]
    total_found: int
    skipped_invalid: int
    skipped_duplicates: int


def extension_of(filename: str) -> str:
    name = (filename or "").strip().lower()
    dot = name.rfind(".")
    return name[dot:] if dot > 0 else ""


def decode_text(data: bytes) -> str:
    """UTF-8, иначе cp1251. BOM отбрасывается."""
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = data.decode("cp1251", errors="replace")
    return text


def _from_text(data: bytes) -> list[str]:
    return _TEXT_EMAIL_RE.findall(decode_text(data))


def _from_xlsx(data: bytes) -> list[str]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover — зависимость объявлена в pyproject
        raise RecipientFileError("На сервере не установлен обработчик XLSX") from exc

    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise RecipientFileError("Не удалось прочитать файл XLSX — он повреждён?") from exc

    found: list[str] = []
    try:
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows(values_only=True):
                for cell in row:
                    if cell is None:
                        continue
                    found.extend(_TEXT_EMAIL_RE.findall(str(cell)))
    finally:
        workbook.close()
    return found


def _from_xls(data: bytes) -> list[str]:
    found: list[str] = []
    try:
        import xlrd

        book = xlrd.open_workbook(file_contents=data)
        for sheet in book.sheets():
            for row_index in range(sheet.nrows):
                for cell in sheet.row_values(row_index):
                    if cell in (None, ""):
                        continue
                    found.extend(_TEXT_EMAIL_RE.findall(str(cell)))
    except Exception as exc:
        # Запасной путь, а не отказ: формат старый, файлы бывают нестандартные,
        # а адреса в них всё равно лежат текстом внутри контейнера.
        logger.info("xls не разобрался библиотекой (%s), ищем адреса в содержимом", exc)

    if not found and data.startswith(_OLE2_MAGIC):
        for encoding in ("latin-1", "utf-16-le"):
            text = data.decode(encoding, errors="ignore")
            found.extend(_BINARY_EMAIL_RE.findall(text))
    return found


def extract_emails(filename: str, data: bytes) -> ImportResult:
    """Разобрать файл и вернуть нормализованные адреса без повторов.

    Порядок сохраняется: он задаёт очерёдность рассылки, и человек, глядя в
    свой файл, ожидает увидеть тот же порядок.
    """
    if not data:
        raise RecipientFileError("Файл пуст")
    if len(data) > MAX_FILE_BYTES:
        raise RecipientFileError(
            f"Файл больше {MAX_FILE_BYTES // (1024 * 1024)} МБ — разделите список на части"
        )

    ext = extension_of(filename)
    if ext not in ALLOWED_EXTENSIONS:
        raise RecipientFileError(
            "Принимаются файлы " + ", ".join(e.upper().lstrip(".") for e in ALLOWED_EXTENSIONS)
        )

    if ext == ".xlsx":
        raw = _from_xlsx(data)
    elif ext == ".xls":
        raw = _from_xls(data)
    else:
        raw = _from_text(data)

    emails: list[str] = []
    seen: set[str] = set()
    invalid = 0
    duplicates = 0
    for candidate in raw:
        value = normalize_email(candidate.strip().strip(".,;"))
        if not is_valid_email(value):
            invalid += 1
            continue
        if value in seen:
            duplicates += 1
            continue
        seen.add(value)
        emails.append(value)

    if not emails:
        raise RecipientFileError("В файле не нашлось ни одного адреса")

    return ImportResult(
        emails=emails,
        total_found=len(raw),
        skipped_invalid=invalid,
        skipped_duplicates=duplicates,
    )
